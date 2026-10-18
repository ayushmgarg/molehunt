"""Write a DOCX analysis report with evidence identity, rankings, and limitations."""

from __future__ import annotations

from collections.abc import Sequence
from datetime import UTC, datetime
from pathlib import Path

from docx import Document
from docx.shared import Inches

from molehunt.detect.score import CandidateScore


def write_analysis_report(
    output_path: str | Path,
    *,
    evidence_name: str,
    evidence_sha256: str,
    observed_slots: int,
    total_slots: int,
    scores: Sequence[CandidateScore],
    flip_rate: float,
    case_id: str = "Not provided",
) -> Path:
    """Create a DOCX report; likelihood ratio uses a binomial null and flip model."""
    if len(evidence_sha256) != 64 or any(
        character not in "0123456789abcdef" for character in evidence_sha256
    ):
        raise ValueError("evidence SHA-256 must be 64 lowercase hexadecimal characters")
    if not 0 <= observed_slots <= total_slots:
        raise ValueError("observed slot count must be between zero and total slots")
    if not 0 < flip_rate < 0.5:
        raise ValueError("flip rate must be strictly between zero and 0.5")
    output = Path(output_path)
    if output.suffix.casefold() != ".docx":
        raise ValueError("analysis reports must use the .docx extension")
    output.parent.mkdir(parents=True, exist_ok=True)

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    document.add_heading("MoleHunt Evidence Analysis", level=0)
    document.add_paragraph(f"Generated UTC: {datetime.now(UTC).isoformat(timespec='seconds')}")
    document.add_heading("Case and evidence", level=1)
    document.add_paragraph(f"Case ID: {case_id}")
    document.add_paragraph(f"Evidence file: {evidence_name}")
    document.add_paragraph(f"Evidence SHA-256: {evidence_sha256}")
    document.add_paragraph(f"Observed slots: {observed_slots} / {total_slots}")

    document.add_heading("Method", level=1)
    document.add_paragraph(
        "Candidate codewords are compared with observed bits. Erased slots are omitted. "
        "The null model uses a 0.5 match probability; the alternative uses the configured "
        f"flip rate e={flip_rate:g}. LR = (1-e)^m e^(n-m) / 0.5^n, where m is matches "
        "and n is observed slots."
    )

    document.add_heading("Candidate ranking", level=1)
    table = document.add_table(rows=1, cols=7)
    table.style = "Light Shading Accent 1"
    headers = ("Rank", "ID", "Matches", "Observed", "Match fraction", "P(null)", "log10 LR")
    for cell, header in zip(table.rows[0].cells, headers, strict=True):
        cell.text = header
    for rank, score in enumerate(scores, start=1):
        row = table.add_row().cells
        values = (
            str(rank),
            score.candidate_id,
            str(score.matches),
            str(score.observed_slots),
            "—" if score.match_fraction is None else f"{score.match_fraction:.3f}",
            f"{score.binomial_p:.4g}",
            f"{score.log10_lr:.3f}",
        )
        for cell, value in zip(row, values, strict=True):
            cell.text = value

    document.add_heading("Interpretation and limitations", level=1)
    document.add_paragraph(
        "This table ranks synthetic copy fingerprints. It does not identify who handled or "
        "disclosed a document. The decision threshold has not been calibrated unless a "
        "separate benchmark establishes its error rate. Results depend on the supplied "
        "master, roster, observed evidence, and flip-rate assumption. Erasures reduce the "
        "number of scored slots."
    )
    document.save(output)
    return output
