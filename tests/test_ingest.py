"""Tests for evidence ingestion and immutable acquisition metadata."""

import hashlib
from datetime import date

import pytest
from docx import Document
from PIL import Image

from molehunt.channels.printdots import DotMetadata, render_dot_grid
from molehunt.detect.ingest import ingest


def test_ingest_utf8_text_preserves_codepoints_and_hash(tmp_path) -> None:
    content = "evidence\u200b café".encode("utf-8")
    path = tmp_path / "leak.txt"
    path.write_bytes(content)
    evidence = ingest(path)
    assert evidence.text == "evidence\u200b café"
    assert evidence.sha256 == hashlib.sha256(content).hexdigest()
    assert evidence.size_bytes == len(content)
    assert evidence.format == "txt"


def test_ingest_docx_paragraphs(tmp_path) -> None:
    path = tmp_path / "leak.docx"
    document = Document()
    document.add_paragraph("First line")
    document.add_paragraph("Second line")
    document.save(path)
    assert ingest(path).text == "First line\nSecond line"


def test_ingest_docx_table_cells(tmp_path) -> None:
    path = tmp_path / "table.docx"
    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Badge"
    table.cell(0, 1).text = "18:42"
    document.save(path)
    assert ingest(path).text == "Badge\t18:42"


def test_ingest_png_reads_dot_metadata_when_configured(tmp_path) -> None:
    path = tmp_path / "scan.png"
    metadata = DotMetadata(7, date(2026, 10, 18), 33)
    render_dot_grid(Image.new("RGB", (160, 100), "white"), metadata, origin=(8, 8), cell_size=6).save(path)
    evidence = ingest(path, dot_origin=(8, 8), dot_cell_size=6)
    assert evidence.text is None
    assert evidence.image_size == (160, 100)
    assert evidence.dot_metadata == metadata


def test_unsupported_and_invalid_files_are_rejected(tmp_path) -> None:
    unsupported = tmp_path / "leak.pdf"
    unsupported.write_bytes(b"data")
    with pytest.raises(ValueError, match="supported"):
        ingest(unsupported)
    invalid = tmp_path / "bad.txt"
    invalid.write_bytes(b"\xff")
    with pytest.raises(ValueError, match="UTF-8"):
        ingest(invalid)


def test_malformed_docx_is_rejected(tmp_path) -> None:
    path = tmp_path / "bad.docx"
    path.write_bytes(b"not a package")
    with pytest.raises(ValueError, match="DOCX"):
        ingest(path)
