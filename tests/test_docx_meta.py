"""Tests for deterministic DOCX metadata fingerprinting."""

import re
from io import BytesIO
from zipfile import ZipFile

from docx import Document
from lxml import etree

from molehunt.channels.docx_meta import derive_rsid, embed_docx_metadata

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
VT = "http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes"


def make_docx() -> bytes:
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("A synthetic ")
    paragraph.add_run("evidence report for the review.")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def test_derived_rsid_is_eight_uppercase_hex_characters() -> None:
    first = derive_rsid(b"k" * 32, "E001")
    assert len(first) == 8
    assert first == derive_rsid(b"k" * 32, "E001")
    assert first != derive_rsid(b"k" * 32, "E002")
    assert all(character in "0123456789ABCDEF" for character in first)


def test_metadata_package_opens_and_preserves_text() -> None:
    result = embed_docx_metadata(make_docx(), b"k" * 32, "E001")
    reopened = Document(BytesIO(result))
    assert "".join(paragraph.text for paragraph in reopened.paragraphs) == (
        "A synthetic evidence report for the review."
    )
    assert len(reopened.paragraphs[0].runs) > 2


def test_settings_runs_and_custom_property_are_fingerprinted() -> None:
    result = embed_docx_metadata(make_docx(), b"k" * 32, "E001")
    with ZipFile(BytesIO(result)) as package:
        settings = etree.fromstring(package.read("word/settings.xml"))
        document = etree.fromstring(package.read("word/document.xml"))
        custom = etree.fromstring(package.read("docProps/custom.xml"))
        relations = etree.fromstring(package.read("_rels/.rels"))

    root_id = settings.get(f"{{{W}}}rsidRoot")
    assert root_id == derive_rsid(b"k" * 32, "E001")
    assert settings.get(f"{{{W}}}rsidRDefault") == root_id
    paragraphs = document.findall(f".//{{{W}}}p")
    runs = document.findall(f".//{{{W}}}r")
    assert paragraphs and all(
        re.fullmatch(r"[0-9A-F]{8}", paragraph.get(f"{{{W}}}rsidP", ""))
        for paragraph in paragraphs
    )
    run_ids = {run.get(f"{{{W}}}rsidR") for run in runs}
    assert None not in run_ids and len(run_ids) >= 2
    properties = custom.findall("{*}property")
    tag = next(prop for prop in properties if prop.get("name") == "MoleHuntTag")
    assert tag.find(f"{{{VT}}}lpwstr").text == "E001"
    assert any(rel.get("Type", "").endswith("/custom-properties") for rel in relations)


def test_reapplying_metadata_updates_instead_of_duplicating_property() -> None:
    once = embed_docx_metadata(make_docx(), b"k" * 32, "E001")
    twice = embed_docx_metadata(once, b"k" * 32, "E002")
    with ZipFile(BytesIO(twice)) as package:
        custom = etree.fromstring(package.read("docProps/custom.xml"))
    tags = [prop for prop in custom.findall("{*}property") if prop.get("name") == "MoleHuntTag"]
    assert len(tags) == 1
    assert tags[0].find(f"{{{VT}}}lpwstr").text == "E002"


def test_metadata_output_is_deterministic_for_same_input() -> None:
    source = make_docx()
    first = embed_docx_metadata(source, b"k" * 32, "E001")
    second = embed_docx_metadata(source, b"k" * 32, "E001")
    assert first == second


def test_missing_or_bad_docx_is_rejected() -> None:
    try:
        embed_docx_metadata(b"not a zip", b"k" * 32, "E001")
    except ValueError:
        pass
    else:
        raise AssertionError("invalid DOCX should be rejected")


def test_rendered_docx_package_is_byte_deterministic() -> None:
    import zipfile
    from io import BytesIO

    from molehunt.core.render import render_docx
    from molehunt.core.slots import allocate_slots
    from molehunt.core.template import parse_template

    template = parse_template("Plan to {{lex:begin|start}} the work.")
    slots = allocate_slots(template, b"k" * 32)
    first = render_docx(template, slots, [1], b"k" * 32, "E001")
    with zipfile.ZipFile(BytesIO(first)) as package:
        assert {info.date_time for info in package.infolist()} == {(2000, 1, 1, 0, 0, 0)}
    assert first == render_docx(template, slots, [1], b"k" * 32, "E001")
