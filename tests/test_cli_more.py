"""Tests for roster, batch issue, collusion analysis, tamper demo, and case reports."""

from pathlib import Path

from docx import Document
from typer.testing import CliRunner

from molehunt.cli import app

runner = CliRunner()
MASTER = Path(__file__).parents[1] / "data" / "masters" / "m2_zoning.md"


def _setup(tmp_path: Path) -> tuple[Path, Path, Path]:
    key, db = tmp_path / "k.key", tmp_path / "r.sqlite3"
    assert runner.invoke(app, ["init", "--key-file", str(key), "--db-file", str(db),
                               "--custody-log", str(tmp_path / "c.log")]).exit_code == 0
    roster = tmp_path / "roster.csv"
    result = runner.invoke(app, ["roster", "--count", "12", "--seed", "4", "-o", str(roster)])
    assert result.exit_code == 0 and "12 synthetic employees" in result.stdout
    registered = runner.invoke(app, ["register", str(MASTER), "--master-id", "M2",
                                     "--key-file", str(key), "--db-file", str(db)])
    assert registered.exit_code == 0, registered.stdout
    return key, db, roster


def test_batch_issue_then_collusion_analysis(tmp_path: Path) -> None:
    key, db, roster = _setup(tmp_path)
    out = tmp_path / "copies"
    issued = runner.invoke(app, [
        "issue", "--master", "M2", "--employees", str(roster), "--format", "txt,docx",
        "--out-dir", str(out), "--key-file", str(key), "--db-file", str(db), "--code", "tardos",
    ])
    assert issued.exit_code == 0, issued.stdout
    assert len(list(out.glob("*.txt"))) == 12 and len(list(out.glob("*.docx"))) == 12
    analyzed = runner.invoke(app, [
        "analyze", str(out / "E003.txt"), "--master", str(MASTER), "--employees", str(roster),
        "--key-file", str(key), "--collusion",
    ])
    assert analyzed.exit_code == 0, analyzed.stdout
    assert "Threshold Z =" in analyzed.stdout
    assert "Accused (family-wise false-accusation target 0.001): E003" in analyzed.stdout


def test_single_mode_analysis_prints_attribution_statement(tmp_path: Path) -> None:
    key, db, roster = _setup(tmp_path)
    copy = tmp_path / "E005.txt"
    assert runner.invoke(app, ["issue", "E005", "--master", "M2", "-o", str(copy),
                               "--key-file", str(key), "--db-file", str(db)]).exit_code == 0
    analyzed = runner.invoke(app, ["analyze", str(copy), "--master", str(MASTER),
                                   "--employees", str(roster), "--key-file", str(key)])
    assert "Finding (attribute)" in analyzed.stdout
    assert "copy issued to E005" in analyzed.stdout


def test_custody_demo_reports_broken_sequence(tmp_path: Path) -> None:
    ledger = tmp_path / "demo" / "custody.log"
    result = runner.invoke(app, ["custody", "demo", "--ledger", str(ledger)])
    assert result.exit_code == 0, result.stdout
    assert "Before tamper: VERIFIED: 8 entries" in result.stdout
    assert "After tamper:  BROKEN at seq 7" in result.stdout
    verify = runner.invoke(app, ["custody", "verify", "--ledger", str(ledger)])
    assert verify.exit_code == 1


def test_case_report_after_custody_intake(tmp_path: Path) -> None:
    key, db, roster = _setup(tmp_path)
    copy = tmp_path / "leak.txt"
    runner.invoke(app, ["issue", "E009", "--master", "M2", "-o", str(copy),
                        "--key-file", str(key), "--db-file", str(db)])
    root = tmp_path / "evidence"
    taken = runner.invoke(app, ["custody", "intake", str(copy), "C-001", "--actor", "EX-01",
                                "--root", str(root), "--db-file", str(db)])
    assert taken.exit_code == 0, taken.stdout
    output = tmp_path / "report.docx"
    made = runner.invoke(app, [
        "report", "--case", "C-001", "-o", str(output), "--master", str(MASTER),
        "--employees", str(roster), "--key-file", str(key), "--examiner", "EX-01",
        "--root", str(root), "--db-file", str(db),
    ])
    assert made.exit_code == 0, made.stdout
    document = Document(output)
    text = "\n".join(p.text for p in document.paragraphs)
    cells = "\n".join(c.text for t in document.tables for r in t.rows for c in r.cells)
    assert "copy issued to E009" in text
    assert "hash ok; chain verified (1)" in cells
    assert "Appendix" in text
