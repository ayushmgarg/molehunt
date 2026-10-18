"""Tests for generated analysis report content."""

import pytest
from docx import Document

from molehunt.detect.fuse import FusedScore
from molehunt.detect.score import score_candidates
from molehunt.report import write_analysis_report


def test_analysis_report_records_hash_rankings_and_limitations(tmp_path) -> None:
    scores = score_candidates([1, 0, 1], {"E001": [1, 0, 1], "E002": [0, 0, 0]})
    output = tmp_path / "analysis.docx"
    write_analysis_report(
        output,
        evidence_name="leak.txt",
        evidence_sha256="a" * 64,
        observed_slots=3,
        total_slots=4,
        scores=scores,
        flip_rate=0.15,
    )

    document = Document(output)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "leak.txt" in text
    assert "a" * 64 in text
    assert "threshold has not been calibrated" in text.lower()
    assert any(cell.text == "E001" for row in document.tables[0].rows for cell in row.cells)


def test_analysis_report_can_be_written_as_pdf(tmp_path) -> None:
    output = tmp_path / "analysis.pdf"
    scores = score_candidates([1, 0], {"E001": [1, 0]})
    write_analysis_report(
        output,
        evidence_name="leak.txt",
        evidence_sha256="b" * 64,
        observed_slots=2,
        total_slots=2,
        scores=scores,
        flip_rate=0.15,
    )
    assert output.read_bytes().startswith(b"%PDF-")


def test_analysis_report_includes_log_fusion_explanation(tmp_path) -> None:
    scores = score_candidates([1, 0], {"E001": [1, 0], "E002": [0, 0]})
    fusion = (
        FusedScore("E001", 0.8, 1.2, 0.95, ("print_job \u00d72",), ("direct access", "print_job \u00d72")),
        FusedScore("E002", 0.2, 0.1, 0.05, (), ("indirect access (0.05 base weight)",)),
    )
    output = tmp_path / "fusion.docx"
    write_analysis_report(
        output,
        evidence_name="leak.txt",
        evidence_sha256="c" * 64,
        observed_slots=2,
        total_slots=2,
        scores=scores,
        flip_rate=0.15,
        fused_scores=fusion,
    )
    document = Document(output)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    table_text = "\n".join(
        cell.text for table in document.tables for row in table.rows for cell in row.cells
    )
    assert "Log fusion" in text
    assert "demo weights are assumptions" in text
    assert "print_job \u00d72" in table_text
    assert "0.950000" in table_text


def test_report_rejects_fusion_with_different_candidates(tmp_path) -> None:
    scores = score_candidates([1], {"E001": [1]})
    with pytest.raises(ValueError, match="same candidates"):
        write_analysis_report(
            tmp_path / "bad.docx",
            evidence_name="leak.txt",
            evidence_sha256="d" * 64,
            observed_slots=1,
            total_slots=1,
            scores=scores,
            flip_rate=0.15,
            fused_scores=(FusedScore("E002", 1, 1, 1, (), ()),),
        )


def test_full_report_has_every_section_eight_part(tmp_path) -> None:
    from molehunt.detect.decide import decide
    from molehunt.report import EvidenceItem, ReportInput, write_report

    codewords = {"E001": (0, 1) * 30, "E002": (1, 0) * 30, "E003": (1, 1, 0, 0) * 15}
    scores = score_candidates(list(codewords["E001"]), codewords)
    data = ReportInput(
        evidence=(EvidenceItem("leak.txt", "e" * 64, "2025-03-03T10:00:00Z", "verified"),),
        scores=scores, observed_slots=60, total_slots=80, flip_rate=0.15,
        case_id="C-001", examiner="EX-01", channels={"lex": (30, 30), "zw": (0, 20)},
        decision=decide(scores, candidates=3),
        timeline=(("2025-03-03T01:50:00Z", "E001", "badge_entry"),),
        slot_matrix=(("E001", "o", "o", "-"), ("E002", "x", "x", "-")),
        hex_excerpt=("0000  41 E2 80 8B  A...",),
        bench_dir=tmp_path / "no-bench",
    )
    output = write_report(tmp_path / "full.docx", data)
    document = Document(output)
    headings = [p.text for p in document.paragraphs if p.style.name.startswith("Heading")]
    assert headings == [
        "Case and scope", "Evidence", "Method", "Findings", "Log fusion",
        "Interpretation and limitations", "Conclusion", "Appendix",
    ]
    text = "\n".join(p.text for p in document.paragraphs)
    assert "copy issued to E001" in text
    assert "Erasure rate: 25.0%" in text
    assert "error rate unknown" in text
    assert "MoleHunt 0." in text
    pdf = write_report(tmp_path / "full.pdf", data)
    assert pdf.read_bytes().startswith(b"%PDF-")
